"""Getting programmes in: AcadDoc course files, and an Excel workbook.

The workbook is for peer universities, whose syllabi usually arrive as PDFs:
someone copies the course list and topics into the template, one course per
row. `write_template` makes a blank one.
"""
from __future__ import annotations

import json
import re
from io import BytesIO
from pathlib import Path
from typing import BinaryIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError

from .schema import Course, Programme

PROGRAMME_FIELDS = [("Institution", "institution"), ("Programme", "name"),
                    ("Source", "source"), ("Syllabus year", "year")]
COURSE_COLUMNS = ["Code", "Title", "Semester", "Credits", "Category", "Topics", "Outcomes"]


class ImportError_(ValueError):
    """An input file that can't be turned into a programme; message says where."""


# ----- AcadDoc -----------------------------------------------------------

def course_from_acaddoc(data: dict) -> Course:
    topics: list[str] = []
    for module in data.get("modules", []):
        topics.append(module["title"])
        topics.extend(p["text"] for p in module.get("parts", []))
    for exercise in data.get("exercises", []):
        topics.append(exercise["title"])
        for item in exercise.get("items", []):
            topics.append(item["text"])
            topics.extend(item.get("subitems", []))
    offerings = data.get("offerings") or [{}]
    return Course(
        code=data["course_code"],
        title=data["course_title"],
        credits=data.get("hours", {}).get("credits"),
        semester=offerings[0].get("semester"),
        category=data.get("category", ""),
        topics=[t for t in topics if t.strip()],
        outcomes=[o["text"] for o in data.get("outcomes", [])],
    )


def from_acaddoc(paths: list[Path], institution: str, name: str,
                 programme_code: str | None = None) -> Programme:
    """Build a programme from AcadDoc course JSON files (or folders of them).

    With `programme_code`, keep only courses offered to that programme
    (AcadDoc's `offerings[].programmes`).
    """
    files: list[Path] = []
    for p in paths:
        files.extend(sorted(p.glob("*.json")) if p.is_dir() else [p])
    return from_acaddoc_texts([(f.name, f.read_text(encoding="utf-8")) for f in files],
                              institution, name, programme_code)


def from_acaddoc_texts(files: list[tuple[str, str]], institution: str, name: str,
                       programme_code: str | None = None) -> Programme:
    """As `from_acaddoc`, from (file name, JSON text) pairs, e.g. browser uploads."""
    courses = []
    for fname, text in files:
        try:
            data = json.loads(text)
            if programme_code and not any(programme_code in o.get("programmes", [])
                                          for o in data.get("offerings", [])):
                continue
            courses.append(course_from_acaddoc(data))
        except (KeyError, TypeError, AttributeError, json.JSONDecodeError, ValidationError) as exc:
            raise ImportError_(f"{fname}: not an AcadDoc course file ({exc})") from exc
    if not courses:
        raise ImportError_("no courses found" + (f" for programme {programme_code}" if programme_code else ""))
    return Programme(institution=institution, name=name, source="AcadDoc", courses=courses)


# ----- Excel workbook ----------------------------------------------------

def _split(cell) -> list[str]:
    if cell is None:
        return []
    return [part.strip(" \t-•") for part in re.split(r"[\n;]", str(cell)) if part.strip(" \t-•")]


def from_workbook(source: str | Path | bytes | BinaryIO, filename: str | None = None) -> Programme:
    """Read the template workbook from a path, bytes or a binary file object."""
    if isinstance(source, (str, Path)):
        filename = filename or Path(source).name
    elif isinstance(source, bytes):
        source = BytesIO(source)
    fname = filename or "workbook"
    try:
        wb = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types
        raise ImportError_(f"{fname}: can't open as an Excel workbook ({exc})") from exc
    for sheet in ("Programme", "Courses"):
        if sheet not in wb.sheetnames:
            raise ImportError_(f"{fname}: no sheet named {sheet!r} (use `curintel template`)")

    meta = {}
    labels = dict(PROGRAMME_FIELDS)
    for row in wb["Programme"].iter_rows(values_only=True):
        if row and row[0] in labels and len(row) > 1 and row[1] is not None:
            meta[labels[row[0]]] = str(row[1]).strip()

    rows = wb["Courses"].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
    missing = [c for c in ("Code", "Title") if c not in header]
    if missing:
        raise ImportError_(f"{fname}: Courses sheet has no {', '.join(missing)} column")
    col = {name: header.index(name) for name in COURSE_COLUMNS if name in header}

    def get(row, name):
        i = col.get(name)
        return row[i] if i is not None and i < len(row) else None

    courses = []
    for n, row in enumerate(rows, start=2):
        if not any(v not in (None, "") for v in row):
            continue
        try:
            credits = get(row, "Credits")
            courses.append(Course(
                code=str(get(row, "Code") or "").strip(),
                title=str(get(row, "Title") or "").strip(),
                semester=None if get(row, "Semester") is None else str(get(row, "Semester")).strip(),
                credits=None if credits in (None, "") else float(credits),
                category=str(get(row, "Category") or "").strip(),
                topics=_split(get(row, "Topics")),
                outcomes=_split(get(row, "Outcomes")),
            ))
        except (ValidationError, ValueError) as exc:
            first = exc.errors()[0] if isinstance(exc, ValidationError) else {"loc": ("Credits",), "msg": str(exc)}
            raise ImportError_(f"{fname}: Courses row {n}, {first['loc'][0]}: {first['msg']}") from exc
    try:
        return Programme(**meta, courses=courses)
    except ValidationError as exc:
        first = exc.errors()[0]
        where = "Courses sheet" if first["loc"][0] == "courses" else f"Programme sheet, {first['loc'][0]}"
        raise ImportError_(f"{fname}: {where}: {first['msg']}") from exc


def write_template(path: str | Path | BinaryIO, programme: Programme | None = None) -> None:
    """A blank template, or (with `programme`) one filled in for review and editing."""
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2F4A6D")

    ws = wb.active
    ws.title = "Programme"
    ws.append(["Field", "Value"])
    for label, attr in PROGRAMME_FIELDS:
        ws.append([label, getattr(programme, attr) if programme else None])
    ws.append([])
    ws.append(["Institution and Programme are required. Source: the URL or document the syllabus came from."])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 60

    cs = wb.create_sheet("Courses")
    cs.append(COURSE_COLUMNS)
    for c in (programme.courses if programme else []):
        cs.append([c.code, c.title, c.semester, c.credits, c.category,
                   "\n".join(c.topics), "\n".join(c.outcomes)])
        for cell in cs[cs.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [10, 34, 10, 8, 12, 70, 50]
    for i, w in enumerate(widths):
        cs.column_dimensions[chr(65 + i)].width = w
    for sheet in (ws, cs):
        for c in sheet[1]:
            c.font, c.fill = head, fill
        sheet.freeze_panes = "A2"

    notes = wb.create_sheet("How to fill")
    for line in [
        "One row per course in the Courses sheet, starting at row 2 (example at the bottom of this sheet).",
        "Topics: one syllabus line per line in the cell (Alt+Enter), or separate them with ;",
        "Copy the topic lines as written in the syllabus - the analysis looks for subject terms in them.",
        "Outcomes are optional. Credits must be a number.",
    ]:
        notes.append([line])
    notes.append([])
    notes.append(COURSE_COLUMNS)
    notes.append(["MTH101", "Calculus", "I", 4, "Core",
                  "Limits and continuity\nDifferentiation and its applications\nRiemann integration",
                  "Evaluate limits and derivatives; Apply integration to areas and volumes"])
    for c in notes[notes.max_row]:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    notes.column_dimensions["A"].width = 100
    wb.save(path)
