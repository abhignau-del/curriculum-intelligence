"""Outcome attainment: how far each course outcome (CO), and through the CO-PO mapping each
programme outcome (PO), was attained by a batch of students.

Input is aggregated, never student by student: for each course outcome and assessment, how
many students were assessed on it and how many reached the target mark. Assessments are of
three kinds - ``internal`` and ``external`` (direct) and ``survey`` (indirect, e.g. a course
exit survey: how many agreed the outcome was achieved).

The policy (YAML, see ``POLICY_DEFAULTS``) turns those counts into levels:

* the share of students reaching the target gives a level by ``thresholds``
  (default: 70% or more -> 3, 60% -> 2, 50% -> 1, below -> 0);
* direct attainment is the weighted mean of the internal and external levels
  (default 30:70); overall attainment weighs direct against survey (default 80:20).
  When a kind of assessment is missing for an outcome, the weights of the others are
  rescaled, and the result says so;
* an outcome is *attained* when its overall value reaches ``target_level`` (default 2).

PO attainment is the strength-weighted mean of the attainment of the course outcomes that
serve the PO (strengths 1-3 from ``mapping``): sum(attainment x strength) / sum(strength).

The defaults are common practice in Indian higher education, not a regulation. Institutions
set their own, and the report prints the policy it used.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal, Optional

import yaml
from pydantic import BaseModel, ConfigDict, Field, ValidationError, model_validator

from .mapping import MappingReport

Kind = Literal["internal", "external", "survey"]
DIRECT: tuple[Kind, ...] = ("internal", "external")

POLICY_DEFAULTS = {
    "name": "Default policy (common practice; set your institution's own)",
    "thresholds": [[0.7, 3], [0.6, 2], [0.5, 1]],
    "weights": {"internal": 0.3, "external": 0.7},
    "direct_share": 0.8,
    "target_level": 2.0,
}


class AttainmentError(ValueError):
    pass


def _message(exc: ValidationError) -> str:
    """The first problem, readably: "thresholds: ..." rather than ": Value error, thresholds: ..."."""
    err = exc.errors()[0]
    where = ".".join(map(str, err["loc"]))
    msg = err["msg"].removeprefix("Value error, ")
    return f"{where}: {msg}" if where else msg


class Policy(BaseModel):
    model_config = ConfigDict(extra="forbid")
    name: str = POLICY_DEFAULTS["name"]
    thresholds: list[tuple[float, int]] = Field(default_factory=lambda: [tuple(t) for t in POLICY_DEFAULTS["thresholds"]])
    weights: dict[str, float] = Field(default_factory=lambda: dict(POLICY_DEFAULTS["weights"]))
    direct_share: float = Field(default=POLICY_DEFAULTS["direct_share"], ge=0, le=1)
    target_level: float = Field(default=POLICY_DEFAULTS["target_level"], ge=0, le=3)

    @model_validator(mode="after")
    def _check(self) -> "Policy":
        if not self.thresholds:
            raise ValueError("thresholds: at least one is needed")
        for share, level in self.thresholds:
            if not 0 <= share <= 1 or level not in (1, 2, 3):
                raise ValueError("thresholds: each is [share 0-1, level 1-3]")
        if set(self.weights) - set(DIRECT):
            raise ValueError("weights: only internal and external")
        if any(w < 0 for w in self.weights.values()) or not sum(self.weights.values()):
            raise ValueError("weights: not negative, and not all zero")
        self.thresholds = sorted(self.thresholds, key=lambda t: -t[0])
        return self

    def level(self, share: float) -> int:
        return next((lv for cut, lv in self.thresholds if share >= cut - 1e-9), 0)

    @classmethod
    def load(cls, path: str | Path) -> "Policy":
        path = Path(path)
        try:
            return cls.model_validate(yaml.safe_load(path.read_text(encoding="utf-8")) or {})
        except ValidationError as exc:
            raise AttainmentError(f"{path.name}: {_message(exc)}") from None


class Record(BaseModel):
    """One assessment of one course outcome, for a batch."""
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)
    course: str = Field(min_length=1)
    co: int = Field(ge=1)                   # 1-based number of the outcome in the course's list
    kind: Kind
    students: int = Field(ge=1)
    reached: int = Field(ge=0)
    note: str = ""

    @model_validator(mode="after")
    def _check(self) -> "Record":
        if self.reached > self.students:
            raise ValueError("reached is more than students")
        return self


class AttainmentData(BaseModel):
    model_config = ConfigDict(extra="forbid")
    batch: str = ""                         # e.g. "2023-26" or "2024-25, semester V"
    records: list[Record] = Field(min_length=1)


@dataclass
class COAttainment:
    course: str
    course_title: str
    co: int
    text: str
    levels: dict[str, int]                  # kind -> level 0-3, for the kinds with data
    shares: dict[str, float]                # kind -> share of students reaching the target
    direct: Optional[float]
    overall: Optional[float]
    notes: list[str] = field(default_factory=list)

    @property
    def label(self) -> str:
        return f"{self.course} CO{self.co}"


@dataclass
class AttainmentReport:
    data: AttainmentData
    policy: Policy
    mapping: MappingReport
    cos: list[COAttainment]
    po_values: list[tuple[str, str, Optional[float], int]]    # (code, name, attainment, outcomes counted)
    unknown: list[str]                      # records naming a course or CO the programme lacks
    no_data: list[str]                      # course outcomes with no record

    def attained(self) -> list[COAttainment]:
        return [c for c in self.cos if c.overall is not None and c.overall >= self.policy.target_level - 1e-9]

    @property
    def assessed(self) -> list[COAttainment]:
        return [c for c in self.cos if c.overall is not None]


def _weighted(values: dict[str, float], weights: dict[str, float]) -> tuple[Optional[float], bool]:
    """Weighted mean over the kinds present; True when weights had to be rescaled."""
    present = {k: w for k, w in weights.items() if k in values and w > 0}
    if not present:
        return None, False
    total = sum(present.values())
    return sum(values[k] * w for k, w in present.items()) / total, len(present) < len([w for w in weights.values() if w > 0])


def compute(mapping: MappingReport, data: AttainmentData, policy: Optional[Policy] = None) -> AttainmentReport:
    policy = policy or Policy()
    programme = mapping.programme
    by_course = {c.code: c for c in programme.courses}
    grouped: dict[tuple[str, int], dict[str, list[Record]]] = {}
    unknown = []
    for r in data.records:
        course = by_course.get(r.course)
        if course is None or r.co > len(course.outcomes):
            unknown.append(f"{r.course} CO{r.co}" + ("" if course else " (no such course)"))
            continue
        grouped.setdefault((r.course, r.co), {}).setdefault(r.kind, []).append(r)

    cos, no_data = [], []
    for c in programme.courses:
        for i, text in enumerate(c.outcomes, start=1):
            kinds = grouped.get((c.code, i))
            if not kinds:
                no_data.append(f"{c.code} CO{i}")
                cos.append(COAttainment(c.code, c.title, i, text, {}, {}, None, None))
                continue
            shares = {k: sum(r.reached for r in rs) / sum(r.students for r in rs) for k, rs in kinds.items()}
            levels = {k: policy.level(s) for k, s in shares.items()}
            notes = []
            direct, rescaled = _weighted({k: float(v) for k, v in levels.items() if k in DIRECT}, policy.weights)
            if rescaled:
                notes.append("only " + " and ".join(k for k in DIRECT if k in levels) + " assessment")
            if direct is None and "survey" not in levels:
                overall = None
            elif direct is None:
                overall = float(levels["survey"])
                notes.append("survey only")
            elif "survey" in levels:
                overall = policy.direct_share * direct + (1 - policy.direct_share) * levels["survey"]
            else:
                overall = direct
                if policy.direct_share < 1:
                    notes.append("no survey")
            cos.append(COAttainment(c.code, c.title, i, text, levels, shares, direct, overall, notes))

    # PO attainment through the CO-PO links. The mapping lists each course's outcomes in order,
    # so the n-th mapped outcome of a course is its CO n.
    att = {(a.course, a.co): a for a in cos}
    numbered, seen = [], {}
    for co in mapping.cos:
        seen[co.course_code] = seen.get(co.course_code, 0) + 1
        numbered.append((seen[co.course_code], co))
    po_values = []
    for po in mapping.framework.outcomes:
        num = den = 0.0
        counted = 0
        for n, co in numbered:
            a = att.get((co.course_code, n))
            link = next((li for li in co.links if li.po == po.code), None)
            if link and a and a.overall is not None:
                num += a.overall * link.strength
                den += link.strength
                counted += 1
        po_values.append((po.code, po.name, num / den if den else None, counted))
    return AttainmentReport(data, policy, mapping, cos, po_values, unknown, no_data)


# ----- reading and writing attainment data -------------------------------------------

SHEET = "Attainment"
COLUMNS = ["Course", "Course title", "CO", "Outcome", "Kind", "Students", "Reached target", "Note"]


def write_template(path, programme, batch: str = "") -> None:
    """A workbook to fill in: one row per course outcome and kind of assessment."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(COLUMNS)
    for c in programme.courses:
        for i, text in enumerate(c.outcomes, start=1):
            for kind in ("internal", "external", "survey"):
                ws.append([c.code, c.title, i, text, kind, None, None, None])
    for i, w in enumerate([12, 26, 5, 60, 10, 10, 15, 30]):
        ws.column_dimensions[chr(65 + i)].width = w
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F4A6D")
    for row in ws.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    meta = wb.create_sheet("Batch")
    meta.append(["Batch", batch])
    meta.column_dimensions["A"].width = 12
    meta.column_dimensions["B"].width = 40
    notes = wb.create_sheet("How to fill")
    for line in [
        "One row per course outcome (CO) and kind of assessment. Fill in Students and Reached target; leave a row blank "
        "if that outcome was not assessed that way.",
        "Students: how many students were assessed on the outcome. Reached target: how many of them reached the target "
        "mark set by your policy (for example 60% of the marks on the questions for that outcome).",
        "Kinds: internal (continuous / internal assessment), external (university examination), survey (course exit "
        "survey: how many agreed the outcome was achieved).",
        "Counts only - no student names or marks are needed. The Batch sheet names the batch, e.g. 2023-26.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    wb.save(path)


def read_data(source, filename: Optional[str] = None) -> AttainmentData:
    """Attainment data from the workbook above, or from JSON ({batch, records}).

    `source` is a path, or the bytes of an uploaded file (then `filename` gives its type).
    """
    import json
    from io import BytesIO

    if isinstance(source, (bytes, bytearray)):
        name = filename or "upload"
        raw, opener = bytes(source), BytesIO(source)
    else:
        path = Path(source)
        name = filename or path.name
        raw, opener = None, path

    try:
        if name.lower().endswith(".json"):
            text = raw.decode("utf-8") if raw is not None else Path(source).read_text(encoding="utf-8")
            return AttainmentData.model_validate(json.loads(text))
    except ValidationError as exc:
        raise AttainmentError(f"{name}: {_message(exc)}") from None
    except (json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise AttainmentError(f"{name}: not valid JSON ({exc})") from None
    from openpyxl import load_workbook

    try:
        wb = load_workbook(opener, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types
        raise AttainmentError(f"{name}: can't open as an Excel workbook ({exc})") from exc
    if SHEET not in wb.sheetnames:
        raise AttainmentError(f"{name}: no sheet named {SHEET!r} (download the entry sheet first)")
    rows = wb[SHEET].iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(rows, [])]
    missing = [c for c in ("Course", "CO", "Kind", "Students", "Reached target") if c not in head]
    if missing:
        raise AttainmentError(f"{name}: {SHEET} sheet has no {', '.join(missing)} column")
    col = {name: head.index(name) for name in COLUMNS if name in head}
    records = []
    for n, row in enumerate(rows, start=2):
        get = lambda name: row[col[name]] if name in col and col[name] < len(row) else None  # noqa: E731
        if get("Students") in (None, "") and get("Reached target") in (None, ""):
            continue                         # not assessed this way
        try:
            records.append(Record(course=str(get("Course") or ""), co=int(get("CO") or 0),
                                  kind=str(get("Kind") or "").strip().lower(),
                                  students=int(get("Students") or 0), reached=int(get("Reached target") or 0),
                                  note=str(get("Note") or "")))
        except (ValidationError, ValueError, TypeError) as exc:
            msg = (_message(exc) if isinstance(exc, ValidationError)
                   else "Students and Reached target must be whole numbers")
            raise AttainmentError(f"{name}: {SHEET} row {n}: {msg}") from None
    batch = ""
    if "Batch" in wb.sheetnames:
        for row in wb["Batch"].iter_rows(values_only=True):
            if row and row[0] == "Batch" and len(row) > 1 and row[1]:
                batch = str(row[1]).strip()
    if not records:
        raise AttainmentError(f"{name}: no outcome has Students and Reached target filled in")
    return AttainmentData(batch=batch, records=records)
