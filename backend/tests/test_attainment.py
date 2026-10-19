"""Outcome attainment: the policy, CO levels, PO roll-up through the mapping, inputs and reports."""
import json

import pytest
from openpyxl import load_workbook

from curintel.attainment import (AttainmentData, AttainmentError, Policy, Record, compute, read_data,
                                 write_template)
from curintel.cli import main
from curintel.mapping import Framework, map_outcomes
from curintel.report import attainment_dict, policy_words, section_attainment
from curintel.schema import Course, Programme
from curintel.taxonomy import Taxonomy
from conftest import SAMPLES

TAX = Taxonomy.builtin()
FW = Framework.from_dict({"id": "t", "name": "T", "outcomes": [
    {"code": "P1", "name": "Solving", "verbs": ["solve"]},                       # strength 1 from a verb
    {"code": "P2", "name": "Proof", "verbs": ["prove"], "words": ["theorem"]},   # 2 when verb and word match
    {"code": "P3", "name": "Ethics", "words": ["ethics"]},
]})
PROG = Programme(institution="U", name="B", courses=[
    Course(code="C1", title="Analysis", outcomes=["Solve limits", "Prove the mean value theorem"]),
    Course(code="C2", title="Algebra", outcomes=["Solve linear systems"]),
    Course(code="C3", title="Seminar", outcomes=[]),
])


def rec(course, co, kind, students, reached):
    return Record(course=course, co=co, kind=kind, students=students, reached=reached)


# ----- the policy ------------------------------------------------------------------

@pytest.mark.parametrize("share,level", [(1.0, 3), (0.7, 3), (0.6999, 2), (0.6, 2), (0.5, 1), (0.4999, 0), (0, 0)])
def test_levels(share, level):
    assert Policy().level(share) == level


@pytest.mark.parametrize("data,message", [
    ({"thresholds": [[1.2, 3]]}, "share 0-1, level 1-3"),
    ({"thresholds": [[0.5, 4]]}, "share 0-1, level 1-3"),
    ({"thresholds": []}, "at least one"),
    ({"weights": {"internal": 1, "midterm": 1}}, "only internal and external"),
    ({"weights": {"internal": 0, "external": 0}}, "not all zero"),
    ({"direct_share": 1.5}, "less than or equal to 1"),
    ({"colour": "red"}, "Extra inputs"),
])
def test_bad_policies(data, message):
    with pytest.raises(ValueError, match=message):
        Policy.model_validate(data)


def test_policy_file_and_words(tmp_path):
    f = tmp_path / "p.yaml"
    f.write_text("name: Ours\nthresholds: [[0.5, 1], [0.8, 3], [0.65, 2]]\nweights: {internal: 1, external: 1}\n"
                 "direct_share: 1\ntarget_level: 1.5\n")
    pol = Policy.load(f)
    assert pol.thresholds == [(0.8, 3), (0.65, 2), (0.5, 1)] and pol.level(0.7) == 2
    f.write_text("thresholds: [[2, 3]]\n")
    with pytest.raises(AttainmentError, match="p.yaml: thresholds"):
        Policy.load(f)
    rep = compute(map_outcomes(PROG, FW, TAX), AttainmentData(records=[rec("C1", 1, "internal", 10, 9)]), pol)
    assert policy_words(rep) == ("level 3 when 80% or more of students reach the target, 2 at 65%, 1 at 50%, "
                                 "otherwise 0. Direct attainment = internal 50 : external 50. Overall = direct 100 : "
                                 "survey 0. Target level 1.5.")


def test_records_are_checked():
    with pytest.raises(ValueError, match="reached is more than students"):
        rec("C1", 1, "internal", 10, 11)
    with pytest.raises(ValueError):
        rec("C1", 0, "internal", 10, 1)
    with pytest.raises(ValueError):
        rec("C1", 1, "midterm", 10, 1)


# ----- computing ----------------------------------------------------------------------

def worked():
    data = AttainmentData(batch="2023-26", records=[
        # C1 CO1: internal 45 of 60 (75% -> 3), external 36 of 60 (60% -> 2), survey 30 of 60 (50% -> 1)
        rec("C1", 1, "internal", 30, 20), rec("C1", 1, "internal", 30, 25),    # two internal tests, pooled
        rec("C1", 1, "external", 60, 36), rec("C1", 1, "survey", 60, 30),
        # C1 CO2: external only, 20 of 60 (33% -> 0)
        rec("C1", 2, "external", 60, 20),
        # C2 CO1: survey only, 50 of 60 (83% -> 3)
        rec("C2", 1, "survey", 60, 50),
        rec("C9", 1, "internal", 10, 5), rec("C2", 4, "internal", 10, 5),     # no such outcomes
    ])
    return compute(map_outcomes(PROG, FW, TAX), data)


def test_co_attainment_worked_example():
    rep = worked()
    co = {c.label: c for c in rep.cos}
    a = co["C1 CO1"]
    assert a.shares == {"internal": 0.75, "external": 0.6, "survey": 0.5}
    assert a.levels == {"internal": 3, "external": 2, "survey": 1}
    assert a.direct == pytest.approx(0.3 * 3 + 0.7 * 2)                     # 2.3
    assert a.overall == pytest.approx(0.8 * 2.3 + 0.2 * 1)                  # 2.04
    assert a.notes == []
    b = co["C1 CO2"]
    assert (b.direct, b.overall) == (0.0, 0.0) and b.notes == ["only external assessment", "no survey"]
    c = co["C2 CO1"]
    assert (c.direct, c.overall, c.notes) == (None, 3.0, ["survey only"])
    assert [x.label for x in rep.attained()] == ["C1 CO1", "C2 CO1"] and len(rep.assessed) == 3
    assert rep.unknown == ["C9 CO1 (no such course)", "C2 CO4"]
    assert rep.no_data == []


def test_po_attainment_is_strength_weighted():
    rep = worked()
    values = {code: (value, n) for code, _, value, n in rep.po_values}
    # P1 (solve): C1 CO1 overall 2.04 and C2 CO1 overall 3, strength 1 each
    assert values["P1"][0] == pytest.approx((2.04 + 3) / 2) and values["P1"][1] == 2
    # P2 (prove + theorem): only C1 CO2, overall 0
    assert values["P2"] == (0.0, 1)
    assert values["P3"] == (None, 0)                                      # no outcome serves it


def test_strengths_weigh_po_attainment():
    fw = Framework.from_dict({"id": "t", "name": "T", "outcomes": [
        {"code": "P", "verbs": ["solve", "prove"], "words": ["theorem"]}]})
    prog = Programme(institution="U", name="B", courses=[
        Course(code="C1", title="A", outcomes=["Solve limits", "Prove the theorem"])])
    data = AttainmentData(records=[rec("C1", 1, "internal", 10, 9), rec("C1", 2, "internal", 10, 0)])
    rep = compute(map_outcomes(prog, fw, TAX), data, Policy(direct_share=1))
    # CO1 level 3 at strength 1, CO2 level 0 at strength 2 (verb and word): (3*1 + 0*2) / 3
    assert rep.po_values[0][2] == pytest.approx(1.0)


def test_outcomes_without_data():
    rep = compute(map_outcomes(PROG, FW, TAX), AttainmentData(records=[rec("C1", 1, "internal", 10, 9)]))
    assert rep.no_data == ["C1 CO2", "C2 CO1"]
    assert rep.cos[1].overall is None and len(rep.assessed) == 1


# ----- inputs ----------------------------------------------------------------------------

def test_template_round_trip(tmp_path):
    write_template(tmp_path / "a.xlsx", PROG, batch="2024-27")
    wb = load_workbook(tmp_path / "a.xlsx")
    ws = wb["Attainment"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 3 * 3 and rows[0][:5] == ("C1", "Analysis", 1, "Solve limits", "internal")
    ws.cell(row=2, column=6).value, ws.cell(row=2, column=7).value = 60, 45
    ws.cell(row=4, column=6).value, ws.cell(row=4, column=7).value = 60, 50
    wb.save(tmp_path / "a.xlsx")
    data = read_data(tmp_path / "a.xlsx")
    assert data.batch == "2024-27"
    assert [(r.course, r.co, r.kind, r.students, r.reached) for r in data.records] == [
        ("C1", 1, "internal", 60, 45), ("C1", 1, "survey", 60, 50)]


@pytest.mark.parametrize("students,reached,message", [
    (60, 70, "row 2: reached is more than students"),
    ("sixty", 10, "row 2: Students and Reached target must be whole numbers"),
])
def test_bad_rows_are_explained(tmp_path, students, reached, message):
    write_template(tmp_path / "a.xlsx", PROG)
    wb = load_workbook(tmp_path / "a.xlsx")
    wb["Attainment"].cell(row=2, column=6).value = students
    wb["Attainment"].cell(row=2, column=7).value = reached
    wb.save(tmp_path / "a.xlsx")
    with pytest.raises(AttainmentError, match=message):
        read_data(tmp_path / "a.xlsx")


def test_other_input_problems(tmp_path):
    write_template(tmp_path / "empty.xlsx", PROG)
    with pytest.raises(AttainmentError, match="no outcome has Students and Reached target filled in"):
        read_data(tmp_path / "empty.xlsx")
    (tmp_path / "bad.json").write_text("{")
    with pytest.raises(AttainmentError, match="not valid JSON"):
        read_data(tmp_path / "bad.json")
    (tmp_path / "bad2.json").write_text(json.dumps({"records": []}))
    with pytest.raises(AttainmentError, match="records"):
        read_data(tmp_path / "bad2.json")
    from openpyxl import Workbook
    Workbook().save(tmp_path / "other.xlsx")
    with pytest.raises(AttainmentError, match="no sheet named 'Attainment'"):
        read_data(tmp_path / "other.xlsx")


# ----- reports and commands ------------------------------------------------------------------

def test_report_section_and_json():
    rep = worked()
    html = section_attainment(rep)
    assert "<h2>Outcome attainment: 2023-26</h2>" in html
    assert "2 of 3 assessed course outcomes reached the target level of 2." in html
    assert "<span class='pill ok'>Met</span>" in html and "<span class='pill High'>Below</span>" in html
    assert "<span class='pill soft'>No data</span>" in html
    assert "Data for outcomes the programme does not have, left out: C9 CO1 (no such course), C2 CO4." in html
    assert "1 course outcome below the target level" in html
    assert "75% <span class='muted small'>→ 3</span>" in html
    d = attainment_dict(rep)
    assert d["attained"] == 2 and d["pos"][2] == {"code": "P3", "name": "Ethics", "attainment": None, "outcomes": 0}
    assert d["policy"]["target_level"] == 2.0
    json.dumps(d)


def test_sample_batch():
    from curintel.schema import load_programme
    prog = load_programme(SAMPLES / "riverside.json")
    rep = compute(map_outcomes(prog, Framework.builtin(), TAX), read_data(SAMPLES / "riverside-attainment.json"))
    assert len(rep.assessed) == 77 and not rep.unknown and not rep.no_data
    values = {code: value for code, _, value, _ in rep.po_values}
    assert values["GA3"] < 2 <= values["GA4"]           # proof-heavy outcomes were made harder in the sample


def test_cli(tmp_path, capsys):
    assert main(["attainment-template", str(SAMPLES / "riverside.json"), "-o", str(tmp_path / "t.xlsx")]) == 0
    assert "77 course outcomes" in capsys.readouterr().out
    assert main(["attainment", str(SAMPLES / "riverside.json"), str(SAMPLES / "riverside-attainment.json"),
                 "-o", str(tmp_path / "a.html"), "--json", str(tmp_path / "a.json")]) == 0
    out = capsys.readouterr().out
    assert "batch 2023-26 (fictional)" in out and "GA3" in out and "below target" in out
    assert "<h2>Outcome attainment" in (tmp_path / "a.html").read_text(encoding="utf-8")
    assert json.loads((tmp_path / "a.json").read_text())["assessed"] == 77
    (tmp_path / "p.yaml").write_text("target_level: 1\n")
    assert main(["attainment", str(SAMPLES / "riverside.json"), str(SAMPLES / "riverside-attainment.json"),
                 "--policy", str(tmp_path / "p.yaml")]) == 0
    assert "reached level 1" in capsys.readouterr().out
    bare = tmp_path / "bare.json"
    bare.write_text(json.dumps({"institution": "U", "name": "B", "courses": [{"code": "C", "title": "T"}]}))
    assert main(["attainment-template", str(bare)]) == 1
    assert "states no course outcomes" in capsys.readouterr().err


def test_exactly_at_the_target_counts_as_attained():
    data = AttainmentData(records=[rec("C1", 1, "internal", 10, 6), rec("C1", 1, "external", 10, 6)])   # 60% -> 2
    rep = compute(map_outcomes(PROG, FW, TAX), data, Policy(direct_share=1))
    assert rep.cos[0].overall == pytest.approx(2.0) and [c.label for c in rep.attained()] == ["C1 CO1"]
