import json
import os
import subprocess
import sys

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from curintel.api import create_app, seed
from curintel.store import Store
from conftest import SAMPLES
from test_io import ACADDOC_COURSE, _fill


@pytest.fixture
def client(tmp_path):
    store = Store(tmp_path / "t.db")
    seed(store)
    return TestClient(create_app(store))


def ids(client):
    progs = client.get("/api/programmes").json()
    own = next(p["id"] for p in progs if p["institution"].startswith("Riverside"))
    return own, [p["id"] for p in progs if p["id"] != own]


def test_library_listing_and_detail(client):
    progs = client.get("/api/programmes").json()
    assert len(progs) == 9
    assert [p["institution"] for p in progs] == sorted(p["institution"] for p in progs)
    own, _ = ids(client)
    d = client.get(f"/api/programmes/{own}").json()
    assert d["programme"]["institution"] == "Riverside College of Science"
    first = d["programme"]["courses"][0]["code"]
    assert "calculus" in d["course_skills"][first]
    assert d["levels"]["python"] == "absent"
    assert client.get("/api/programmes/nope").status_code == 404


def test_benchmark(client):
    own, peers = ids(client)
    r = client.post("/api/benchmark", json={"own": own, "peers": peers})
    assert r.status_code == 200
    d = r.json()
    assert d["programme"]["id"] == own
    assert {p["id"] for p in d["peers"]} == set(peers)
    py = next(s for s in d["skills"] if s["id"] == "python")
    assert py["status"] == "gap" and py["priority"] == "High"
    # evidence for the click-through: peer label -> syllabus lines
    assert any(e["line"] for evs in py["peer_evidence"].values() for e in evs)
    assert len(d["overlaps"]) == 2


def test_benchmark_rejects_bad_selections(client):
    own, peers = ids(client)
    assert client.post("/api/benchmark", json={"own": own, "peers": []}).status_code == 422
    r = client.post("/api/benchmark", json={"own": own, "peers": [own, *peers]})
    assert r.status_code == 400 and "can't also be" in r.json()["detail"]
    assert client.post("/api/benchmark", json={"own": own, "peers": ["missing"]}).status_code == 404


def test_duplicate_peer_institution_is_a_400(client):
    own, peers = ids(client)
    copy = client.get(f"/api/programmes/{peers[0]}").json()["programme"]
    dup = client.post("/api/programmes", json=copy).json()["id"]
    r = client.post("/api/benchmark", json={"own": own, "peers": [peers[0], dup]})
    assert r.status_code == 400 and "same institution" in r.json()["detail"]


def test_report_download(client):
    own, peers = ids(client)
    r = client.post("/api/benchmark/report", json={"own": own, "peers": peers})
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    assert "Priority gaps" in r.text


def test_upload_json_and_xlsx(client, tmp_path):
    body = (SAMPLES / "peers" / "westgate.json").read_bytes()
    r = client.post("/api/programmes/upload", params={"filename": "w.json"}, content=body)
    assert r.status_code == 201 and r.json()["courses"] > 10

    f = tmp_path / "peer.xlsx"
    _fill(f, [["MA1", "Calculus", "I", 4, "Core", "Limits", None]])
    r = client.post("/api/programmes/upload", params={"filename": "peer.xlsx"}, content=f.read_bytes())
    assert r.status_code == 201 and r.json()["institution"] == "Peer U"

    _fill(f, [["MA1", "Calculus", "I", "four", "Core", "Limits", None]])
    r = client.post("/api/programmes/upload", params={"filename": "peer.xlsx"}, content=f.read_bytes())
    assert r.status_code == 400 and "peer.xlsx: Courses row 2, Credits" in r.json()["detail"]

    r = client.post("/api/programmes/upload", params={"filename": "x.json"}, content=b'{"institution": "X"}')
    assert r.status_code == 400 and r.json()["detail"].startswith("x.json: name")
    r = client.post("/api/programmes/upload", params={"filename": "x.pdf"}, content=b"%PDF")
    assert r.status_code == 400


def test_acaddoc_import(client):
    files = [{"name": "MTH101.json", "text": json.dumps(ACADDOC_COURSE)}]
    r = client.post("/api/programmes/acaddoc", json={"institution": "Inst", "name": "B.Tech", "files": files})
    assert r.status_code == 201 and r.json()["courses"] == 1
    r = client.post("/api/programmes/acaddoc",
                    json={"institution": "Inst", "name": "B.Tech", "files": [{"name": "x.json", "text": "[]"}]})
    assert r.status_code == 400 and "x.json" in r.json()["detail"]


def test_replace_delete_download(client):
    own, _ = ids(client)
    p = client.get(f"/api/programmes/{own}").json()["programme"]
    p["year"] = "2026-27"
    assert client.put(f"/api/programmes/{own}", json=p).json()["year"] == "2026-27"
    r = client.get(f"/api/programmes/{own}/download")
    assert json.loads(r.text)["year"] == "2026-27" and "attachment" in r.headers["content-disposition"]
    assert client.delete(f"/api/programmes/{own}").status_code == 204
    assert client.delete(f"/api/programmes/{own}").status_code == 404
    assert len(client.get("/api/programmes").json()) == 8


def test_template_and_taxonomy(client, tmp_path):
    r = client.get("/api/template.xlsx")
    (tmp_path / "t.xlsx").write_bytes(r.content)
    assert "Courses" in load_workbook(tmp_path / "t.xlsx").sheetnames
    tax = client.get("/api/taxonomy").json()
    assert sum(len(a["skills"]) for a in tax["areas"]) == 48


def test_report_identical_across_hash_seeds(tmp_path):
    """Set iteration order changes with PYTHONHASHSEED; the report must not."""
    outs = []
    for seed_value in ("1", "2"):
        out = tmp_path / f"r{seed_value}.json"
        env = dict(os.environ, PYTHONHASHSEED=seed_value)
        subprocess.run([sys.executable, "-m", "curintel", "report", str(SAMPLES / "riverside.json"),
                        str(SAMPLES / "peers"), "-o", str(tmp_path / "r.html"), "--json", str(out)],
                       check=True, env=env, capture_output=True)
        outs.append(out.read_text())
    assert outs[0] == outs[1]


def test_upload_size_limit(client):
    r = client.post("/api/programmes/upload", params={"filename": "big.json"}, content=b" " * (10 * 1024 * 1024 + 1))
    assert r.status_code == 413


def test_rules_listing_and_checklist_in_benchmark(client):
    rules = client.get("/api/rules").json()
    ugc = next(r for r in rules if r["id"] == "ugc-ccfup-2022")
    assert [v["id"] for v in ugc["variants"]] == ["4-year-research", "4-year", "3-year"]
    own, peers = ids(client)
    d = client.post("/api/benchmark", json={"own": own, "peers": peers}).json()
    assert d["regulatory"] == []                                  # only when asked
    d = client.post("/api/benchmark", json={"own": own, "peers": peers, "rules": ["ugc-ccfup-2022"]}).json()
    reg = d["regulatory"][0]
    assert reg["variant"] == "3-year" and {r["status"] for r in reg["results"]} == {"met", "partly"}
    d = client.post("/api/benchmark", json={"own": own, "peers": peers, "rules": ["ugc-ccfup-2022"],
                                            "degree": "4-year"}).json()
    assert d["regulatory"][0]["variant"] == "4-year" and d["regulatory"][0]["missing_semesters"] == [7, 8]
    html = client.post("/api/benchmark/report", json={"own": own, "peers": peers, "rules": ["ugc-ccfup-2022"]}).text
    assert "Regulatory checklist" in html


@pytest.mark.parametrize("body,message", [
    ({"rules": ["nope"]}, "No built-in rules named 'nope'"),
    ({"rules": ["../../etc/passwd.yaml"]}, "No built-in rules named"),        # no file paths over HTTP
    ({"rules": ["ugc-ccfup-2022"], "degree": "diploma"}, "matches none of"),
])
def test_bad_rules_requests_are_400(client, body, message):
    own, peers = ids(client)
    r = client.post("/api/benchmark", json={"own": own, "peers": peers, **body})
    assert r.status_code == 400 and message in r.json()["detail"]


def test_programme_detail_includes_scheme(client):
    own, _ = ids(client)
    p = client.get(f"/api/programmes/{own}").json()["programme"]
    assert p["degree"] == "3-year" and len(p["scheme"]) == 24


def test_benchmark_includes_outcomes_and_copo(client):
    own, peers = ids(client)
    d = client.post("/api/benchmark", json={"own": own, "peers": peers}).json()
    assert d["outcomes"]["higher_order_share"] > 0
    assert d["copo"]["framework"]["id"] == "ugc-locf-mathematics" and len(d["copo"]["pos"]) == 9
    r = client.post("/api/benchmark", json={"own": own, "peers": peers, "framework": "../secret.yaml"})
    assert r.status_code == 400 and "No built-in outcome framework" in r.json()["detail"]
    r = client.post("/api/benchmark", json={"own": own, "peers": peers, "framework": "own"})
    assert r.status_code == 400 and "states no programme outcomes" in r.json()["detail"]
    html = client.post("/api/benchmark/report", json={"own": own, "peers": peers}).text
    assert "<h2>CO–PO mapping" in html
    # a peer without outcomes under review: neither section
    peer = peers[0]
    d = client.post("/api/benchmark", json={"own": peer, "peers": [own]}).json()
    assert d["outcomes"] is None and d["copo"] is None


def test_attainment_template_and_upload(client, tmp_path):
    from openpyxl import load_workbook
    own, peers = ids(client)
    r = client.get(f"/api/programmes/{own}/attainment-template.xlsx", params={"batch": "2024-27"})
    assert r.status_code == 200 and "attainment.xlsx" in r.headers["content-disposition"]
    (tmp_path / "a.xlsx").write_bytes(r.content)
    wb = load_workbook(tmp_path / "a.xlsx")
    assert wb["Batch"]["B1"].value == "2024-27"
    ws = wb["Attainment"]
    ws.cell(row=2, column=6).value, ws.cell(row=2, column=7).value = 60, 45      # RMA101 CO1 internal, 75%
    wb.save(tmp_path / "a.xlsx")
    body = (tmp_path / "a.xlsx").read_bytes()
    r = client.post(f"/api/programmes/{own}/attainment", params={"filename": "a.xlsx"}, content=body)
    assert r.status_code == 200
    d = r.json()
    assert d["batch"] == "2024-27" and d["assessed"] == 1 and d["cos"][0]["levels"] == {"internal": 3}
    assert d["policy"]["target_level"] == 2.0
    # a policy of the client's own
    r = client.post(f"/api/programmes/{own}/attainment", content=body,
                    params={"filename": "a.xlsx", "policy": json.dumps({"thresholds": [[0.8, 3], [0.7, 2], [0.5, 1]]})})
    assert r.json()["cos"][0]["levels"] == {"internal": 2}
    html = client.post(f"/api/programmes/{own}/attainment", content=body,
                       params={"filename": "a.xlsx", "format": "html"}).text
    assert "<h2>Outcome attainment: 2024-27</h2>" in html
    assert client.get("/api/attainment/policy").json()["weights"] == {"internal": 0.3, "external": 0.7}


@pytest.mark.parametrize("params,body,message", [
    ({"filename": "a.csv"}, b"x", "expected the attainment workbook"),
    ({"filename": "a.json"}, b"{", "not valid JSON"),
    ({"filename": "a.json", "policy": '{"thresholds": [[2, 3]]}'}, b"{}", "Policy: thresholds: each is"),
    ({"filename": "a.json", "framework": "../x.yaml"}, b"{}", "No built-in outcome framework"),
    ({"filename": "a.json"}, json.dumps({"records": [{"course": "X", "co": 1, "kind": "internal", "students": 5,
                                                       "reached": 9}]}).encode(), "reached is more than students"),
])
def test_attainment_upload_errors(client, params, body, message):
    own, _ = ids(client)
    r = client.post(f"/api/programmes/{own}/attainment", params=params, content=body)
    assert r.status_code == 400 and message in r.json()["detail"]


def test_attainment_needs_outcomes(client):
    own, peers = ids(client)
    r = client.get(f"/api/programmes/{peers[0]}/attainment-template.xlsx")
    assert r.status_code == 400 and "no course outcomes" in r.json()["detail"]
