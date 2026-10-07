import json

import pytest
from openpyxl import load_workbook

from curintel.cli import main
from curintel.importers import ImportError_, from_acaddoc, from_workbook, write_template
from curintel.schema import load_programme, save_programme
from conftest import SAMPLES, prog

ACADDOC_COURSE = {
    "course_code": "MTH101", "course_title": "Linear Algebra and Calculus", "kind": "theory",
    "category": "Foundation", "offerings": [{"semester": "I", "programmes": ["CSE", "ME"]}],
    "hours": {"credits": 4}, "outcomes": [{"code": "CO1", "text": "Solve linear systems."}],
    "modules": [{"title": "Matrices", "parts": [{"text": "Rank of a matrix"}, {"text": "Eigenvalues"}]}],
    "exercises": [],
}


def test_acaddoc_import(tmp_path):
    (tmp_path / "MTH101.json").write_text(json.dumps(ACADDOC_COURSE))
    other = dict(ACADDOC_COURSE, course_code="EE1", offerings=[{"semester": "II", "programmes": ["EE"]}])
    (tmp_path / "EE1.json").write_text(json.dumps(other))
    p = from_acaddoc([tmp_path], "Inst", "B.Tech")
    assert len(p.courses) == 2
    c = next(c for c in p.courses if c.code == "MTH101")
    assert c.topics == ["Matrices", "Rank of a matrix", "Eigenvalues"]
    assert c.outcomes == ["Solve linear systems."] and c.credits == 4 and c.semester == "I"
    only = from_acaddoc([tmp_path], "Inst", "B.Tech", programme_code="CSE")
    assert [c.code for c in only.courses] == ["MTH101"]


def test_acaddoc_import_reports_bad_file(tmp_path):
    (tmp_path / "bad.json").write_text('{"course_title": "no code"}')
    with pytest.raises(ImportError_, match="bad.json"):
        from_acaddoc([tmp_path], "I", "P")


def _fill(path, rows, meta=(("Institution", "Peer U"), ("Programme", "B.Sc. Maths"))):
    write_template(path)
    wb = load_workbook(path)
    for i, (k, v) in enumerate(meta, start=2):
        assert wb["Programme"].cell(i, 1).value == k
        wb["Programme"].cell(i, 2).value = v
    for r in rows:
        wb["Courses"].append(r)
    wb.save(path)


def test_workbook_round_trip(tmp_path):
    f = tmp_path / "peer.xlsx"
    _fill(f, [["MA1", "Calculus", "I", 4, "Core", "Limits\nDerivatives; Integrals", "Find limits"],
              [None] * 7,
              ["MA2", "Python", 2, None, None, "- NumPy", None]])
    p = load_programme(f)
    assert p.institution == "Peer U" and p.name == "B.Sc. Maths"
    assert p.courses[0].topics == ["Limits", "Derivatives", "Integrals"]
    assert p.courses[0].outcomes == ["Find limits"]
    assert p.courses[1].semester == "2" and p.courses[1].credits is None
    assert p.courses[1].topics == ["NumPy"]


def test_workbook_errors_name_the_place(tmp_path):
    f = tmp_path / "peer.xlsx"
    _fill(f, [["MA1", "Calculus", "I", "four", "Core", "Limits", None]])
    with pytest.raises(ImportError_, match="row 2, Credits"):
        from_workbook(f)
    _fill(f, [["MA1", None, "I", 4, "Core", "Limits", None]])
    with pytest.raises(ImportError_, match="row 2, title"):
        from_workbook(f)
    _fill(f, [["MA1", "Calculus", "I", 4, "Core", "Limits", None]], meta=(("Institution", None),))
    with pytest.raises(ImportError_, match="Programme sheet, institution"):
        from_workbook(f)
    _fill(f, [])
    with pytest.raises(ImportError_, match="Courses sheet"):
        from_workbook(f)


def test_bad_json_names_the_file(tmp_path):
    f = tmp_path / "x.json"
    f.write_text('{"institution": "X", "name": "Y", "courses": [{"code": "A"}]}')
    with pytest.raises(ValueError, match=r"x.json: courses.0.title"):
        load_programme(f)


def test_cli_report(tmp_path, capsys):
    out, data = tmp_path / "r.html", tmp_path / "r.json"
    assert main(["report", str(SAMPLES / "riverside.json"), str(SAMPLES / "peers"),
                 "-o", str(out), "--json", str(data)]) == 0
    html = out.read_text(encoding="utf-8")
    assert "Priority gaps" in html and "Statistical Methods" in html
    d = json.loads(data.read_text())
    assert len(d["peers"]) == 8
    assert any(s["id"] == "python" and s["status"] == "gap" for s in d["skills"])
    assert "[High  ] Optimisation" in capsys.readouterr().out


def test_cli_reports_errors_without_traceback(tmp_path, capsys):
    bad = tmp_path / "bad.json"
    bad.write_text("{}")
    assert main(["report", str(bad), str(SAMPLES / "peers")]) == 1
    assert "error: bad.json" in capsys.readouterr().err


def test_report_escapes_html(tmp_path, tax):
    from curintel.analysis import benchmark
    from curintel.report import render_html
    own = prog("<script>alert(1)</script>", ("Calculus", ["<b>limits</b>"]))
    html = render_html(benchmark(own, [prog("Peer", ("Linear Programming", []))], tax), [])
    assert "<script>alert" not in html and "&lt;script&gt;" in html


def test_cli_check_and_template(tmp_path, capsys):
    f = tmp_path / "p.json"
    save_programme(prog("X", ("Calculus", ["Arrays and functions"])), f)
    assert main(["check", str(f)]) == 0
    out = capsys.readouterr().out
    assert "Calculus" in out and "C0: Arrays and functions" in out
    assert main(["template", "-o", str(tmp_path / "t.xlsx")]) == 0
    assert set(load_workbook(tmp_path / "t.xlsx").sheetnames) == {"Programme", "Courses", "How to fill"}
