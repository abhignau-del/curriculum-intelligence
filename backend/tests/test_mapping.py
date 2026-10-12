"""CO-PO mapping: frameworks, indicators, strengths, a programme's own outcomes, reports and inputs."""
import json

import pytest
from openpyxl import load_workbook

from curintel.cli import main
from curintel.importers import ImportError_, from_workbook, write_template
from curintel.mapping import Framework, FrameworkError, builtin_frameworks, choose_framework, derive_po, map_outcomes
from curintel.outcomes import Bloom
from curintel.report import copo_dict, render_outcomes_html, section_copo
from curintel.schema import Course, Programme, ProgrammeOutcome, load_programme
from curintel.taxonomy import Taxonomy
from conftest import SAMPLES

TAX = Taxonomy.builtin()
BLOOM = Bloom.builtin()
FW = Framework.from_dict({"id": "t", "name": "Test POs", "source": "src", "outcomes": [
    {"code": "P1", "name": "Knowledge", "skills": ["area:pure"], "level": 2, "via_course": True},
    {"code": "P2", "name": "Problem solving", "verbs": ["solve", "compute"], "words": ["real life"], "level": 3},
    {"code": "P3", "name": "Digital", "skills": ["python"], "words": ["software"]},
    {"code": "P4", "name": "Ethics", "words": ["ethics"]},
]})


def prog(*courses, pos=()):
    """(code, title, topics, outcomes) tuples."""
    return Programme(institution="U", name="B", programme_outcomes=list(pos), courses=[
        Course(code=c, title=t, semester="I", topics=tp, outcomes=o) for c, t, tp, o in courses])


def links(m, text):
    co = next(c for c in m.cos if c.text == text)
    return {li.po: (li.strength, li.reasons) for li in co.links}


def test_indicators_and_strengths():
    m = map_outcomes(prog(("C1", "Linear Algebra", ["Vector spaces"], [
        "Solve real life problems with matrices",           # verb + words: two kinds; Apply reaches level 3
        "Compute eigenvalues of a matrix",                   # verb only, level 3 reached
        "Use Python software for linear systems",            # P3: skill + words
        "Define a vector space",                             # skill named (linear algebra), Remember < level 2
        "Realise the beauty of the subject",                 # nothing named: P1 only through the course
    ])), FW, TAX)
    assert links(m, "Solve real life problems with matrices")["P2"] == (3, ["verb: solve", "words: real life"])
    assert links(m, "Compute eigenvalues of a matrix")["P2"][0] == 2
    assert links(m, "Compute eigenvalues of a matrix")["P1"] == (2, ["skill: Linear Algebra"])
    assert links(m, "Use Python software for linear systems")["P3"] == (2, ["skill: Python", "words: software"])
    assert links(m, "Define a vector space")["P1"] == (1, ["skill: Linear Algebra"])
    assert links(m, "Realise the beauty of the subject") == {"P1": (1, ["course: Linear Algebra"])}
    assert m.unmapped == []


def test_course_matrix_and_po_summary():
    m = map_outcomes(prog(("C1", "Linear Algebra", [], ["Solve real life problems", "Compute ranks"]),
                          ("C2", "History", [], ["Discuss ethics in research"])), FW, TAX)
    matrix = {code: cells for code, _, cells in m.course_matrix()}
    assert matrix["C1"]["P2"] == 3                       # mean of 3 and 2 = 2.5, rounded up
    assert "P1" in matrix["C1"] and matrix["C2"] == {"P4": 1}
    summary = {po.code: (n, mean) for po, n, mean in m.po_summary()}
    assert summary["P2"] == (1, 3.0) and summary["P3"] == (0, None) and summary["P4"] == (1, 1.0)


def test_unmapped_outcomes():
    m = map_outcomes(prog(("C2", "History", [], ["Appreciate the past"])), FW, TAX)
    assert [co.text for co in m.unmapped] == ["Appreciate the past"]


def test_framework_checks():
    with pytest.raises(FrameworkError, match="share a code"):
        Framework.from_dict({"id": "x", "name": "x", "outcomes": [{"code": "A"}, {"code": "A"}]})
    with pytest.raises(FrameworkError, match="missing or invalid"):
        Framework.from_dict({"id": "x", "outcomes": []})
    bad = Framework.from_dict({"id": "x", "name": "x", "outcomes": [{"code": "A", "skills": ["alchemy"]}]})
    with pytest.raises(FrameworkError, match="skill 'alchemy' is not in"):
        map_outcomes(prog(("C", "T", [], ["Solve x"])), bad, TAX)
    bad = Framework.from_dict({"id": "x", "name": "x", "outcomes": [{"code": "A", "skills": ["area:magic"]}]})
    with pytest.raises(FrameworkError, match="area 'magic'"):
        map_outcomes(prog(("C", "T", [], ["Solve x"])), bad, TAX)
    with pytest.raises(FrameworkError, match="no built-in outcome framework named 'nope'"):
        Framework.resolve("nope")


def test_builtin_locf_framework():
    assert "ugc-locf-mathematics" in builtin_frameworks()
    fw = Framework.builtin()
    assert [po.code for po in fw.outcomes] == [f"GA{i}" for i in range(1, 10)]
    assert all(po.page in ("5", "6") for po in fw.outcomes)
    m = map_outcomes(prog(("C", "Graph Theory", [], ["Identify types of graphs and use them to solve problems"])), fw, TAX)
    assert "GA2" not in links(m, "Identify types of graphs and use them to solve problems")   # graphs, not plots


def test_own_outcomes_get_indicators_from_their_wording():
    po = derive_po("PSO1", "", "Apply numerical methods and Python programming to problems in finance", BLOOM, TAX)
    assert po.verbs == ("apply",) and po.level == 3
    assert set(po.skills) == {"numerical_methods", "python", "programming"}
    assert po.words == ("finance",) and po.derived       # "problems", "methods", "python" (a skill) left out
    p = prog(("C1", "Numerics", [], ["Implement Newton-Raphson in Python"]),
             pos=[ProgrammeOutcome(code="PSO1", text="Apply numerical methods and Python programming to real world problems"),
                  ProgrammeOutcome(code="PO2", name="Ethics", text="Act ethically", words=["ethic"], level=1)])
    fw = choose_framework(None, p, TAX)                      # the programme's own, by default
    assert [o.code for o in fw.outcomes] == ["PSO1", "PO2"] and fw.outcomes[1].words == ("ethic",)
    m = map_outcomes(p, fw, TAX)
    assert links(m, "Implement Newton-Raphson in Python")["PSO1"][1] == ["skill: Numerical Methods, Python"]
    assert choose_framework("ugc-locf-mathematics", p, TAX).id == "ugc-locf-mathematics"
    with pytest.raises(FrameworkError, match="states no programme outcomes"):
        choose_framework("own", prog(("C", "T", [], [])), TAX)
    assert choose_framework(None, prog(("C", "T", [], [])), TAX).id == "ugc-locf-mathematics"


def test_report_section_and_json():
    m = map_outcomes(prog(("C1", "Linear <Algebra>", [], ["Solve real life problems"])), FW, TAX)
    html = section_copo(m)
    assert "<h2>CO–PO mapping: Test POs</h2>" in html and "Linear &lt;Algebra&gt;" in html
    assert "<b>P3</b> Digital, <b>P4</b> Ethics are served by no course outcome." in html
    assert "<td class='s3'>3</td>" in html and "Source: src." in html
    d = copo_dict(m)
    assert d["matrix"][0]["cells"]["P2"] == 3 and d["pos"][2] == {"code": "P3", "name": "Digital", "text": "",
                                                                  "page": "", "courses": 0, "mean": None}
    json.dumps(d)
    page = render_outcomes_html(m.outcomes, m)
    assert "<h2>Course outcomes</h2>" in page and "<h2>CO–PO mapping" in page
    empty = section_copo(map_outcomes(prog(("C", "T", [], [])), FW, TAX))
    assert "nothing to map" in empty


def test_sample_programme_against_locf():
    m = map_outcomes(load_programme(SAMPLES / "riverside.json"), Framework.builtin(), TAX)
    summary = {po.code: n for po, n, _ in m.po_summary()}
    assert summary["GA1"] == 21 and summary["GA3"] > 0 and summary["GA8"] == 0
    # History of Mathematics and the project teach no pure, applied or statistics skill
    assert {c for c, _, cells in m.course_matrix() if "GA1" not in cells} == {"RMA602", "RMA603"}


def test_workbook_programme_outcomes(tmp_path):
    p = load_programme(SAMPLES / "riverside.json").model_copy(update={"programme_outcomes": [
        ProgrammeOutcome(code="PSO1", name="Modelling", text="Model real problems mathematically")]})
    write_template(tmp_path / "r.xlsx", p)
    assert from_workbook(tmp_path / "r.xlsx").programme_outcomes == p.programme_outcomes
    wb = load_workbook(tmp_path / "r.xlsx")
    wb["Programme outcomes"].cell(row=2, column=3).value = None
    wb.save(tmp_path / "bad.xlsx")
    with pytest.raises(ImportError_, match="Programme outcomes row 2: needs a Code and an Outcome"):
        from_workbook(tmp_path / "bad.xlsx")


def test_cli_copo_and_report(tmp_path, capsys):
    assert main(["copo", str(SAMPLES / "riverside.json"), "-o", str(tmp_path / "c.html"),
                 "--json", str(tmp_path / "c.json")]) == 0
    out = capsys.readouterr().out
    assert "against UGC LOCF Mathematics graduate attributes" in out and "GA1" in out
    assert "<h2>CO–PO mapping" in (tmp_path / "c.html").read_text(encoding="utf-8")
    assert json.loads((tmp_path / "c.json").read_text())["framework"]["id"] == "ugc-locf-mathematics"
    assert main(["report", str(SAMPLES / "riverside.json"), str(SAMPLES / "peers"), "-o", str(tmp_path / "r.html")]) == 0
    assert "<h2>CO–PO mapping" in (tmp_path / "r.html").read_text(encoding="utf-8")
    assert main(["copo", str(SAMPLES / "riverside.json"), "--framework", "own"]) == 1
    assert "states no programme outcomes" in capsys.readouterr().err


def test_edges_of_the_rules():
    m = map_outcomes(prog(("C1", "Linear Algebra", [], ["Explain the beauty of the subject"]),
                          ("C2", "Python Programming", [], ["Write loops and functions"]),
                          ("C3", "Philosophy", [], ["Discuss metaethics and its history"])), FW, TAX)
    # through the course only: slight, even at the PO's Bloom level (Explain = 2 = P1's level)
    assert links(m, "Explain the beauty of the subject") == {"P1": (1, ["course: Linear Algebra"])}
    # only POs marked via_course are served through the course: P3 (Python) is not
    assert "P3" not in links(m, "Write loops and functions")
    # words match whole words: "ethics" is not in "metaethics"
    assert "P4" not in links(m, "Discuss metaethics and its history")


def test_derived_pos_are_held_to_more():
    pos = [ProgrammeOutcome(code="PO1", text="Use laboratory equipment and work independently"),
           ProgrammeOutcome(code="PO2", text="Analyse and solve well-defined problems"),
           ProgrammeOutcome(code="PO3", text="Develop research temperament through theory and practical learning")]
    texts = ["Use Python to plot graphs", "Solve recurrence relations", "Understand basic concepts of number theory",
             "Design a research project on graphs", "Prove the rank-nullity theorem"]
    texts += [f"Compute example {i} of the methods" for i in range(6)]          # "methods": in over a fifth
    p = prog(("C1", "Discrete Mathematics", [], texts), pos=pos)
    m = map_outcomes(p, choose_framework(None, p, TAX), TAX)
    assert links(m, "Use Python to plot graphs") == {}                  # generic verb alone: no link
    assert links(m, "Solve recurrence relations") == {"PO2": (1, ["verb: solve"])}   # specific verb; Apply < Analyse: no bonus
    assert "PO3" not in links(m, "Understand basic concepts of number theory")       # "theory" is too common a word
    assert links(m, "Design a research project on graphs")["PO3"][1] == ["words: research"]


def test_words_common_in_the_programme_do_not_count_for_derived_pos():
    texts = [f"Compute sequences of type {i}" for i in range(4)] + [f"Define object {i}" for i in range(8)]
    p = prog(("C1", "Analysis", [], texts),
             pos=[ProgrammeOutcome(code="PO1", text="Study convergent sequences carefully")])
    m = map_outcomes(p, choose_framework(None, p, TAX), TAX)
    # "sequences" is in 4 of 12 outcomes (over a fifth): too common to show a link
    assert links(m, "Compute sequences of type 0") == {}
    # the same word, written into a curated framework, still counts
    fw = Framework.from_dict({"id": "x", "name": "x", "outcomes": [{"code": "PO1", "words": ["sequences"]}]})
    assert links(map_outcomes(p, fw, TAX), "Compute sequences of type 0") == {"PO1": (1, ["words: sequences"])}
