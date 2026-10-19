"""The NAAC 2.6.1 evidence pack: what the Word document says, with and without attainment."""
from io import BytesIO

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from fastapi.testclient import TestClient

from curintel.api import create_app, seed
from curintel.attainment import AttainmentData, Policy, Record, compute, read_data
from curintel.cli import main
from curintel.evidence import build, to_bytes
from curintel.mapping import Framework, choose_framework, map_outcomes
from curintel.schema import Course, Programme, ProgrammeOutcome, load_programme
from curintel.store import Store
from curintel.taxonomy import Taxonomy
from conftest import SAMPLES

TAX = Taxonomy.builtin()


def text_of(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def headings(doc) -> list[str]:
    return [p.text for p in doc.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]


def sample(with_attainment=True):
    prog = load_programme(SAMPLES / "riverside.json")
    m = map_outcomes(prog, Framework.builtin(), TAX)
    att = compute(m, read_data(SAMPLES / "riverside-attainment.json")) if with_attainment else None
    return build(m, att), m, att


def test_sections_in_order():
    doc, _, _ = sample()
    hs = headings(doc)
    assert hs[0] == "B.Sc. Mathematics (Honours): learning outcomes and their attainment"
    top = [h for h in hs if h[:2] in ("Be", "1.", "2.", "3.", "4.", "5.", "6.")]
    assert top == ["Before submission", "1. Programme outcomes", "2. Course outcomes", "3. CO–PO articulation matrix",
                   "4. Attainment: batch 2023-26 (fictional)", "5. Actions for outcomes below target",
                   "6. How this was prepared"]
    assert "RMA101 Calculus (semester I)" in hs                     # one heading per course


def test_before_submission_lists_what_the_data_shows():
    doc, m, att = sample()
    text = text_of(doc)
    assert "The programme states no programme outcomes (POs/PSOs) of its own" in text
    assert "16 of 77 course outcomes open with a verb that is hard to assess" in text
    assert "No course outcome serves GA7 Self-directed learning, GA8 Moral and ethical awareness/reasoning, GA9" in text
    assert "22 course outcome(s) fell below the target level" in text
    assert "Add where the outcomes are published" in text
    assert "NAAC metric 2.6.1" in text


def test_course_outcomes_and_matrix():
    doc, m, _ = sample()
    co_table = next(t for t in doc.tables if t.rows[0].cells[0].text == "CO" and t.rows[0].cells[-1].text == "Bloom level")
    assert [c.text for c in co_table.rows[1].cells] == ["CO1", "Understand the basic ideas of Calculus", "2 Understand"]
    matrix = next(t for t in doc.tables if t.rows[0].cells[0].text == "Course")
    assert [c.text for c in matrix.rows[0].cells] == ["Course", *[f"GA{i}" for i in range(1, 10)]]
    first = {code: cells for code, _, cells in m.course_matrix()}["RMA101"]
    row = [c.text for c in matrix.rows[1].cells]
    assert row[0] == "RMA101 Calculus" and row[4] == str(first["GA4"]) and row[5] == ""
    assert matrix.rows[-1].cells[0].text == "Courses serving"
    shaded = matrix.rows[1].cells[4]._tc.xpath(".//w:shd/@w:fill")
    assert shaded == ["3F6E9C"]                                      # strength 3: the darkest shade
    # nine POs: the matrix gets a landscape page; the rest stay portrait, all A4
    orient = [s.orientation for s in doc.sections]
    assert orient == [WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT]
    assert round(doc.sections[0].page_width.inches, 2) == 8.27


def test_attainment_and_actions():
    doc, _, att = sample()
    text = text_of(doc)
    assert "Policy: Default policy (common practice; set your institution's own). Level 3 when 70%" in text
    assert "55 of 77 assessed course outcomes reached the target level of 2." in text
    assert "GA3 Critical thinking and analytical reasoning | 1.27 | 14 | below" in text
    actions = next(t for t in doc.tables if [c.text for c in t.rows[0].cells][2:3] == ["Action taken"])
    below = [c for c in att.assessed if c.overall < 2]
    assert len(actions.rows) == 1 + len(below)
    assert actions.rows[1].cells[0].text == min(below, key=lambda c: c.overall).label   # weakest first
    assert all(r.cells[2].text == "" for r in actions.rows[1:])      # for the department to fill in


def test_without_attainment():
    doc, _, _ = sample(with_attainment=False)
    text = text_of(doc)
    assert "No attainment results were given: section 4 is empty." in text
    assert "No attainment results were given for this draft." in text
    assert "4. Attainment" in headings(doc)
    actions = next(t for t in doc.tables if [c.text for c in t.rows[0].cells][2:3] == ["Action taken"])
    assert len(actions.rows) == 2                                    # one blank row to fill in


def test_own_programme_outcomes_and_few_pos():
    prog = Programme(institution="U", name="B.Sc.", courses=[
        Course(code="C1", title="Algebra", semester="I", outcomes=["Solve linear systems", "Prove Lagrange's theorem"]),
        Course(code="C2", title="Seminar", semester="II")],
        programme_outcomes=[ProgrammeOutcome(code="PO1", name="Problem solving", text="Solve problems in algebra")])
    m = map_outcomes(prog, choose_framework(None, prog, TAX), TAX)
    doc = build(m, compute(m, AttainmentData(records=[Record(course="C1", co=1, kind="internal", students=10, reached=9)]),
                           Policy(direct_share=1)))
    text = text_of(doc)
    assert "The programme's stated programme outcomes and programme-specific outcomes." in text
    assert "states no programme outcomes" not in text
    assert "1 course(s) state no course outcomes: C2 Seminar." in text
    assert "1 course outcome(s) were not assessed in this batch." in text
    assert "PO1 | Problem solving: Solve problems in algebra" in text
    assert [s.orientation for s in doc.sections] == [WD_ORIENT.PORTRAIT] * 3     # one PO: no landscape page
    assert to_bytes(doc)[:2] == b"PK"                                # a real .docx (zip) file


def test_cli(tmp_path, capsys):
    out = tmp_path / "e.docx"
    assert main(["evidence", str(SAMPLES / "riverside.json"), "--attainment", str(SAMPLES / "riverside-attainment.json"),
                 "-o", str(out)]) == 0
    assert "Evidence pack written" in capsys.readouterr().out
    assert "55 of 77" in text_of(Document(out))
    assert main(["evidence", str(SAMPLES / "riverside.json"), "-o", str(out)]) == 0
    assert "(no attainment results given)" in capsys.readouterr().out
    bare = tmp_path / "bare.json"
    bare.write_text('{"institution": "U", "name": "B", "courses": [{"code": "C", "title": "T"}]}')
    assert main(["evidence", str(bare)]) == 1
    assert "no course outcomes" in capsys.readouterr().err


@pytest.fixture
def client(tmp_path):
    store = Store(tmp_path / "t.db")
    seed(store)
    return TestClient(create_app(store))


def test_api(client):
    progs = client.get("/api/programmes").json()
    own = next(p["id"] for p in progs if p["institution"].startswith("Riverside"))
    peer = next(p["id"] for p in progs if p["id"] != own)
    r = client.post(f"/api/programmes/{own}/evidence.docx")
    assert r.status_code == 200 and "NAAC-2.6.1-evidence.docx" in r.headers["content-disposition"]
    assert "No attainment results were given" in text_of(Document(BytesIO(r.content)))
    body = (SAMPLES / "riverside-attainment.json").read_bytes()
    r = client.post(f"/api/programmes/{own}/evidence.docx", params={"filename": "a.json"}, content=body)
    assert "55 of 77" in text_of(Document(BytesIO(r.content)))
    r = client.post(f"/api/programmes/{own}/evidence.docx", params={"filename": "a.csv"}, content=b"x")
    assert r.status_code == 400 and "expected the attainment workbook" in r.json()["detail"]
    r = client.post(f"/api/programmes/{peer}/evidence.docx")
    assert r.status_code == 400 and "no course outcomes" in r.json()["detail"]
    r = client.post(f"/api/programmes/{own}/evidence.docx", params={"policy": '{"target_level": 9}'})
    assert r.status_code == 400 and r.json()["detail"].startswith("Policy: target_level")
